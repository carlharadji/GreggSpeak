import argparse
import csv
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCX = PROJECT_ROOT.parent / "manus" / "week15v5.docx"

HEADERS = [
    "Ground Truth Word",
    "Predicted Word",
    "Classification Result",
    "Stroke Similarity Observation",
    "Possible Cause of Misclassification",
]

V1_ROWS = [
    {
        "true": "act",
        "pred": "contact",
        "observation": "Similar compact loop and rising terminal stroke direction",
        "cause": "Baseline feature extraction overemphasized the shared loop-and-slant pattern, causing confusion with the longer contact outline.",
    },
    {
        "true": "already",
        "pred": "corrected",
        "observation": "Shared flowing horizontal stroke and upward terminal extension",
        "cause": "V1 relied on dominant silhouette-level features and did not separate the finer internal stroke differences between the two outlines.",
    },
    {
        "true": "comply",
        "pred": "complaint",
        "observation": "Similar rounded initial stroke and compact ending loop",
        "cause": "The baseline model confused a shorter outline with a longer legal-word outline because both retained similar curvature and terminal structure.",
    },
    {
        "true": "court",
        "pred": "caught",
        "observation": "Both outlines have compact angled strokes with a similar overall silhouette",
        "cause": "Small curvature differences were not strongly separated in the V1 feature space, resulting in a high-confidence wrong prediction.",
    },
    {
        "true": "failed",
        "pred": "child",
        "observation": "Shared looped starting form and sweeping terminal curve",
        "cause": "The baseline model treated visually close loop-and-tail patterns as the same class-level feature.",
    },
]

V2_ROWS = [
    {
        "true": "debt",
        "pred": "judge",
        "result": "Misclassified",
        "observation": "Presence of similar loop curvature and initial stroke orientation, with comparable compact symbol structure",
        "cause": "High structural similarity causing feature overlap in convolutional layers, leading to confusion between closely related shorthand patterns",
    },
    {
        "true": "doctor",
        "pred": "during",
        "result": "Misclassified",
        "observation": "Similar stroke flow and directional curvature, with nearly identical mid-stroke transitions",
        "cause": "Model reliance on dominant stroke patterns rather than fine-grained distinctions, resulting in misclassification of visually overlapping symbols",
    },
    {
        "true": "here",
        "pred": "her",
        "result": "Misclassified",
        "observation": "Nearly identical base structure with only slight variation in terminal stroke length and extension",
        "cause": "Minimal inter-class variation not sufficiently captured during training, causing the model to treat both patterns as the same feature representation",
    },
]

V3_ROWS = [
    {
        "true": "year",
        "pred": "were",
        "result": "Became misclassified",
        "observation": "Nearly identical shorthand outline and stroke direction, with only small variation in length and curve placement",
        "cause": "The expanded V3 dataset altered the learned feature boundary for this pair, causing year to be classified as were despite improvement in other classes.",
    },
    {
        "true": "her",
        "pred": "here",
        "result": "Still Misclassified",
        "observation": "Nearly identical base structure with only slight variation in terminal stroke length and extension",
        "cause": "Minimal inter-class variation continued to challenge feature separation, causing the model to treat both patterns as the same representation.",
    },
]


def read_incorrect_v1() -> set[tuple[str, str]]:
    path = PROJECT_ROOT / "reports" / "incorrect_predictions_v1.csv"
    pairs: set[tuple[str, str]] = set()
    with path.open(encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            pairs.add((row["true_label"], row["predicted_label"]))
    return pairs


def verify_v1_rows() -> None:
    pairs = read_incorrect_v1()
    missing = [(row["true"], row["pred"]) for row in V1_ROWS if (row["true"], row["pred"]) not in pairs]
    if missing:
        raise RuntimeError(f"Selected V1 rows were not found in incorrect_predictions_v1.csv: {missing}")


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_paragraph_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def find_paragraph(doc: Document, exact_text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise RuntimeError(f"Paragraph not found: {exact_text}")


def replace_paragraph_starting(doc: Document, prefix: str, replacement: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_paragraph_text(paragraph, replacement)
            return
    raise RuntimeError(f"Paragraph starting with not found: {prefix}")


def insert_table_after(paragraph, doc: Document, rows: int, cols: int, style):
    table = doc.add_table(rows=rows, cols=cols)
    if style is not None:
        table.style = style
    paragraph._p.addnext(table._tbl)
    return table


def image_for(label: str, use_testing_crop: bool) -> Path:
    if use_testing_crop:
        path = PROJECT_ROOT / "datasets" / "all_words_testing_uniform" / f"{label}_000.png"
        if path.exists():
            return path
        if label == "its":
            alt = PROJECT_ROOT / "datasets" / "all_words_testing_uniform" / "it’s_000.png"
            if alt.exists():
                return alt
    folder = PROJECT_ROOT / "datasets" / "all_words" / label
    direct = folder / f"{label}_000.png"
    if direct.exists():
        return direct
    matches = sorted(folder.glob(f"{label}_*.png"))
    if matches:
        return matches[0]
    raise FileNotFoundError(f"No image found for label: {label}")


def set_cell_text(cell, text: str, bold: bool = False, font_size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_word_image_cell(cell, label: str, image_path: Path) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(label)
    run.bold = True
    run.font.size = Pt(8)
    run.add_break()
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(0.68))
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def fill_misclassification_table(table, rows: list[dict], font_size: int = 8) -> None:
    for col_idx, header in enumerate(HEADERS):
        set_cell_text(table.cell(0, col_idx), header, bold=True, font_size=font_size)

    for row_idx, row_data in enumerate(rows, start=1):
        true_label = row_data["true"]
        pred_label = row_data["pred"]
        set_word_image_cell(table.cell(row_idx, 0), true_label, image_for(true_label, use_testing_crop=True))
        set_word_image_cell(table.cell(row_idx, 1), pred_label, image_for(pred_label, use_testing_crop=False))
        set_cell_text(table.cell(row_idx, 2), row_data.get("result", "Misclassified"), font_size=font_size)
        set_cell_text(table.cell(row_idx, 3), row_data["observation"], font_size=font_size)
        set_cell_text(table.cell(row_idx, 4), row_data["cause"], font_size=font_size)

    widths = [0.88, 0.88, 0.95, 1.85, 1.95]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def resize_table(table, desired_rows: int, desired_cols: int) -> None:
    while len(table.rows) < desired_rows:
        table.add_row()
    while len(table.rows) > desired_rows:
        table._tbl.remove(table.rows[-1]._tr)
    if len(table.columns) != desired_cols:
        raise RuntimeError(f"Expected {desired_cols} columns, found {len(table.columns)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args()

    verify_v1_rows()

    docx_path = args.docx.resolve()
    if not docx_path.exists():
        raise FileNotFoundError(docx_path)
    if not args.no_backup:
        backup = docx_path.with_name(f"{docx_path.stem}.before_v1_misclass_update{docx_path.suffix}")
        shutil.copy2(docx_path, backup)

    doc = Document(docx_path)
    if len(doc.tables) <= 16:
        raise RuntimeError(f"Expected at least 17 tables, found {len(doc.tables)}")

    v2_table = doc.tables[15]
    v3_table = doc.tables[16]
    table_style = v2_table.style

    replace_paragraph_starting(
        doc,
        "Table 7 presents the results of the internal testing conducted on two versions",
        "Table 7 presents the results of the internal testing conducted on three versions of the shorthand recognition model. The purpose of this evaluation is to observe the progression of the model’s performance as improvements were introduced during the development process.",
    )
    replace_paragraph_starting(
        doc,
        "Despite the strong performance of V1",
        "Because V1 served as the baseline model, it was also examined against the later internal comparison set used for the succeeding model checks. In this comparison, the baseline model produced more misclassifications than the later versions, particularly on shorthand outlines with compact loops, similar stroke direction, or elongated terminal strokes. Representative V1 errors included “act” predicted as “contact,” “already” predicted as “corrected,” “comply” predicted as “complaint,” “court” predicted as “caught,” and “failed” predicted as “child.” These errors show that the baseline model tended to rely on dominant silhouette-level patterns and was less stable in separating fine-grained shorthand variations.",
    )
    replace_paragraph_starting(
        doc,
        "However, despite the improved performance of V2",
        "However, despite the improved performance of V2, a small number of misclassified shorthand words were still observed. Similar to V1, most classification errors involved shorthand symbols with highly overlapping visual structures. Words such as “debt” and “judge,” “doctor” and “during,” as well as “here” and “her,” continued to exhibit structural similarities that caused confusion within the convolutional layers of the model. These results suggest that while V2 improved the system’s ability to recognize shorthand patterns, certain classes still lacked sufficient visual distinction for the model to consistently separate them during prediction.",
    )

    placeholder = find_paragraph(doc, "<<Put the Model v1 missclassified table here>>")
    set_paragraph_text(placeholder, "Table 8. Misclassified Gregg Shorthand Words of Model V1")
    v1_table = insert_table_after(placeholder, doc, rows=len(V1_ROWS) + 1, cols=len(HEADERS), style=table_style)
    fill_misclassification_table(v1_table, V1_ROWS)

    matching_v2_captions = [
        p
        for p in doc.paragraphs
        if p.text.strip() == "Table 8. Misclassified Gregg Shorthand Words of Model V1"
        and p._p is not placeholder._p
    ]
    if not matching_v2_captions:
        raise RuntimeError("Could not find the old mislabeled V2 caption after inserting the V1 table.")
    caption_v2 = matching_v2_captions[0]
    set_paragraph_text(caption_v2, "Table 9. Misclassified Gregg Shorthand Words of Model V2")
    resize_table(v2_table, len(V2_ROWS) + 1, len(HEADERS))
    fill_misclassification_table(v2_table, V2_ROWS)

    caption_v3 = find_paragraph(doc, "Table 10. Misclassified Gregg Shorthand Words of Model V2")
    set_paragraph_text(caption_v3, "Table 10. Misclassified Gregg Shorthand Words of Model V3")
    resize_table(v3_table, len(V3_ROWS) + 1, len(HEADERS))
    fill_misclassification_table(v3_table, V3_ROWS)

    doc.save(docx_path)
    print(f"Updated {docx_path}")


if __name__ == "__main__":
    main()
